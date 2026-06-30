"""New format parser + structural converter tests (other-file-formats work)."""
from __future__ import annotations

import json
from pathlib import Path

import pytest

from dataset_insight.ingestion.jsonl_parser import JsonlParser
from dataset_insight.ingestion.markdown_parser import MarkdownParser
from dataset_insight.ingestion.xlsx_parser import XlsxParser
from dataset_insight.ingestion.docx_parser import DocxParser
from dataset_insight.conversion.rule_based import convert_csv_rows
from dataset_insight.conversion.structural import (
    convert_structural,
    StructuralExtractionError,
)
from dataset_insight.ingestion.base import ParsedDocument


# --- JSONL -----------------------------------------------------------------
def test_jsonl_flat_rows(tmp_path: Path):
    f = tmp_path / "d.jsonl"
    f.write_text(
        '{"instruction": "Q1", "output": "A1"}\n{"instruction": "Q2", "output": "A2"}\n',
        encoding="utf-8",
    )
    parsed = JsonlParser().parse(str(f))
    assert len(parsed.rows) == 2
    assert parsed.rows[0] == {"instruction": "Q1", "output": "A1"}


def test_jsonl_chat_messages_flattened(tmp_path: Path):
    f = tmp_path / "d.jsonl"
    item = {
        "messages": [
            {"role": "system", "content": "sys"},
            {"role": "user", "content": "Hello?"},
            {"role": "assistant", "content": "Hi!"},
        ]
    }
    f.write_text(json.dumps(item, ensure_ascii=False) + "\n", encoding="utf-8")
    parsed = JsonlParser().parse(str(f))
    assert parsed.rows[0] == {"instruction": "Hello?", "output": "Hi!"}


def test_jsonl_sharegpt_conversations(tmp_path: Path):
    f = tmp_path / "d.json"
    item = {
        "conversations": [
            {"from": "human", "value": "Question?"},
            {"from": "gpt", "value": "Answer."},
        ]
    }
    f.write_text(json.dumps([item], ensure_ascii=False), encoding="utf-8")
    parsed = JsonlParser().parse(str(f))
    assert parsed.rows[0] == {"instruction": "Question?", "output": "Answer."}
    # Can flow straight into the pipeline
    pairs = convert_csv_rows(parsed.rows, "instruction", "output")
    assert len(pairs) == 1


# --- Markdown / structural -------------------------------------------------
def test_markdown_sections_to_pairs(tmp_path: Path):
    f = tmp_path / "d.md"
    f.write_text(
        "# What is Python?\nPython is a versatile programming language.\n\n"
        "# Lists\nLists use square brackets, e.g. [1, 2, 3].\n",
        encoding="utf-8",
    )
    parsed = MarkdownParser().parse(str(f))
    assert parsed.sections and len(parsed.sections) == 2
    pairs = convert_structural(parsed)
    assert len(pairs) == 2
    # A heading that is a question stays as-is
    assert pairs[0]["instruction"] == "What is Python?"
    # A non-question heading is wrapped into a phrasing
    assert pairs[1]["instruction"].startswith("Explain")


def test_structural_faq_markers():
    parsed = ParsedDocument(
        source_format="pdf",
        raw_text="Question: What is the capital?\nAnswer: It is Ankara.\n"
        "Question: Population?\nAnswer: Large.",
    )
    pairs = convert_structural(parsed)
    assert len(pairs) == 2
    assert pairs[0]["instruction"] == "What is the capital?"


def test_structural_question_lines():
    parsed = ParsedDocument(
        source_format="pdf",
        raw_text="What is Python?\nA high-level language.\nWhat is it for?\nMany things.",
    )
    pairs = convert_structural(parsed)
    assert len(pairs) == 2


def test_structural_flat_prose_blocks():
    # No structure -> error + LLM recommendation
    parsed = ParsedDocument(
        source_format="pdf",
        raw_text=(
            "This is a long prose text. There are no headings or questions. "
            "Just flowing sentences that keep going."
        ),
    )
    with pytest.raises(StructuralExtractionError) as exc:
        convert_structural(parsed)
    assert "LLM" in str(exc.value)


def test_structural_empty_pdf_scanned():
    parsed = ParsedDocument(source_format="pdf", raw_text="")
    with pytest.raises(StructuralExtractionError) as exc:
        convert_structural(parsed)
    assert "OCR" in str(exc.value)


# --- XLSX / DOCX (create real files) ---------------------------------------
def test_xlsx_roundtrip(tmp_path: Path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["question", "answer", "topic"])
    ws.append(["Q1", "A1", "T1"])
    ws.append(["Q2", "A2", "T2"])
    p = tmp_path / "d.xlsx"
    wb.save(str(p))

    parsed = XlsxParser().parse(str(p))
    assert len(parsed.rows) == 2
    assert parsed.rows[0]["question"] == "Q1"
    pairs = convert_csv_rows(parsed.rows, "question", "answer", "topic")
    assert pairs[0]["category"] == "T1"


def test_docx_headings_to_sections(tmp_path: Path):
    from docx import Document

    doc = Document()
    doc.add_heading("Heading One", level=1)
    doc.add_paragraph("This is the body text of the first section, long enough.")
    doc.add_heading("Heading Two", level=1)
    doc.add_paragraph("The body of the second section goes here too.")
    p = tmp_path / "d.docx"
    doc.save(str(p))

    parsed = DocxParser().parse(str(p))
    assert parsed.sections and len(parsed.sections) == 2
    pairs = convert_structural(parsed)
    assert len(pairs) == 2
